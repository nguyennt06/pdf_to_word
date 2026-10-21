import unittest
from pathlib import Path
from tempfile import TemporaryDirectory

import pymupdf
from docx import Document
from docx.oxml import OxmlElement
from docx.oxml.ns import qn

from main import (
    _find_contents_page_groups,
    _restore_pdf_spaces,
    normalize_docx_xml,
    repair_table_of_contents,
)


def _char(text, x0, x1):
    return {"c": text, "bbox": (x0, 0.0, x1, 12.0), "origin": (x0, 10.0)}


def _text_span(*chars, size=12.0):
    return {"chars": list(chars), "size": size, "style": []}


def _text_line(*spans, direction=(1.0, 0.0)):
    return {
        "dir": direction,
        "wmode": 0,
        "spans": list(spans),
    }


class PdfSpaceRepairTests(unittest.TestCase):
    def test_preserves_whitespace_spans_and_infers_missing_word_space(self):
        line = _text_line(
            _text_span(_char("a", 0.0, 4.0)),
            _text_span(_char(" ", 4.0, 7.0)),
            _text_span(_char("b", 7.0, 11.0), _char("c", 14.0, 18.0)),
        )
        blocks = [{"type": 0, "lines": [line]}]

        self.assertEqual(_restore_pdf_spaces(blocks), (1, 1))
        self.assertEqual(line["spans"][0]["chars"][0]["c"], "a ")
        self.assertEqual(line["spans"][1]["chars"], [])
        self.assertEqual(line["spans"][2]["chars"][0]["c"], "b ")

        # A second pass does not duplicate either repaired space.
        self.assertEqual(_restore_pdf_spaces(blocks), (0, 0))

    def test_ignores_tight_kerning_and_rotated_text(self):
        horizontal = _text_line(
            _text_span(
                _char("A", 0.0, 6.0),
                _char("V", 6.1, 12.0),
            )
        )
        rotated = _text_line(
            _text_span(
                _char("a", 0.0, 4.0),
                _char("b", 8.0, 12.0),
            ),
            direction=(0.0, 1.0),
        )

        self.assertEqual(
            _restore_pdf_spaces([{"type": 0, "lines": [horizontal, rotated]}]),
            (0, 0),
        )
        self.assertEqual(horizontal["spans"][0]["chars"][0]["c"], "A")
        self.assertEqual(rotated["spans"][0]["chars"][0]["c"], "a")

    def test_separates_adjacent_raw_lines_on_the_same_baseline(self):
        first = _text_line(_text_span(_char("a", 0.0, 4.0)))
        second = _text_line(_text_span(_char("b", 9.0, 13.0)))
        blocks = [{"type": 0, "lines": [first, second]}]

        self.assertEqual(_restore_pdf_spaces(blocks), (0, 1))
        self.assertEqual(first["spans"][0]["chars"][0]["c"], "a ")

    def test_separates_words_wrapped_across_pdf_lines(self):
        first_char = _char("n", 0.0, 4.0)
        second_char = _char("d", 0.0, 4.0)
        second_char["origin"] = (0.0, 24.0)
        first = _text_line(_text_span(first_char))
        second = _text_line(_text_span(second_char))

        self.assertEqual(
            _restore_pdf_spaces([{"type": 0, "lines": [first, second]}]),
            (0, 1),
        )
        self.assertEqual(first["spans"][0]["chars"][0]["c"], "n ")


class DocxNormalizationTests(unittest.TestCase):
    def test_normalizes_values_and_reorders_properties(self):
        document = Document()
        paragraph = document.add_paragraph("sample")
        properties = paragraph._p.get_or_add_pPr()

        for name in (
            "autoSpaceDN",
            "tabs",
            "autoSpaceDE",
            "widowControl",
            "spacing",
            "ind",
            "jc",
        ):
            properties.append(OxmlElement(f"w:{name}"))

        scaling = OxmlElement("w:w")
        scaling.set(qn("w:val"), "97.136")
        run_properties = paragraph.runs[0]._r.get_or_add_rPr()
        run_properties.append(scaling)
        run_properties.append(OxmlElement("w:rFonts"))
        run_properties.append(OxmlElement("w:color"))
        run_properties.append(OxmlElement("w:sz"))

        table = document.add_table(rows=1, cols=1)
        indent = OxmlElement("w:tblInd")
        indent.set(qn("w:w"), "180.5")
        table._tbl.tblPr.append(indent)

        borders = OxmlElement("w:tcBorders")
        for name in ("start", "top", "end", "bottom"):
            side = OxmlElement(f"w:{name}")
            side.set(qn("w:color"), "#000000")
            borders.append(side)
        border = borders.find(qn("w:bottom"))
        border.set(qn("w:sz"), "3.184")
        cell_properties = table.cell(0, 0)._tc.get_or_add_tcPr()
        cell_properties.append(borders)
        for side in ("start", "end"):
            margin = OxmlElement("w:tcMar")
            margin.append(OxmlElement(f"w:{side}"))
            cell_properties.append(margin)
        shading = OxmlElement("w:shd")
        shading.set(qn("w:fill"), "EEEEEE")
        cell_properties.append(shading)

        self.assertEqual(normalize_docx_xml(document), 8)
        self.assertEqual(scaling.get(qn("w:val")), "97")
        self.assertEqual(indent.get(qn("w:w")), "181")
        self.assertEqual(border.get(qn("w:sz")), "3")
        self.assertTrue(
            all(side.get(qn("w:color")) == "000000" for side in borders)
        )
        self.assertEqual(shading.get(qn("w:val")), "clear")
        self.assertEqual(
            [child.tag.rsplit("}", 1)[-1] for child in properties],
            [
                "widowControl",
                "tabs",
                "autoSpaceDE",
                "autoSpaceDN",
                "spacing",
                "ind",
                "jc",
            ],
        )
        self.assertEqual(
            [child.tag.rsplit("}", 1)[-1] for child in borders],
            ["top", "start", "bottom", "end"],
        )
        margins = cell_properties.findall(qn("w:tcMar"))
        self.assertEqual(len(margins), 1)
        self.assertEqual(
            [child.tag.rsplit("}", 1)[-1] for child in margins[0]],
            ["start", "end"],
        )
        self.assertLess(
            [child.tag.rsplit("}", 1)[-1] for child in run_properties].index("w"),
            [child.tag.rsplit("}", 1)[-1] for child in run_properties].index("sz"),
        )


class TableOfContentsRepairTests(unittest.TestCase):
    def _make_contents_pdf(self, path: Path) -> None:
        pdf = pymupdf.open()
        first = pdf.new_page(width=595, height=842)
        first.insert_text((220, 55), "TABLE OF CONTENTS", fontsize=14)
        first.insert_text((80, 120), "1.1", fontsize=12)
        first.insert_text((115, 120), "Overview of the whole system", fontsize=12)
        first.insert_text((520, 120), "3", fontsize=12)
        first.insert_text((100, 150), "1.1.1", fontsize=12)
        first.insert_text((140, 150), "Wrapped title start", fontsize=12)
        first.insert_text((140, 166), "continued title", fontsize=12)
        first.insert_text((520, 166), "4", fontsize=12)
        first.insert_text((80, 200), "CONCLUSION", fontsize=12)
        first.insert_text((520, 200), "5", fontsize=12)

        second = pdf.new_page(width=595, height=842)
        second.insert_text((80, 100), "2.1", fontsize=12)
        second.insert_text((115, 100), "Second page section", fontsize=12)
        second.insert_text((520, 100), "8", fontsize=12)
        second.insert_text((100, 130), "2.1.1", fontsize=12)
        second.insert_text((140, 130), "Nested section", fontsize=12)
        second.insert_text((520, 130), "9", fontsize=12)

        third = pdf.new_page(width=595, height=842)
        third.insert_text((80, 100), "Abbreviations", fontsize=12)
        pdf.save(path)
        pdf.close()

    @staticmethod
    def _section_marker():
        paragraph = OxmlElement("w:p")
        properties = OxmlElement("w:pPr")
        properties.append(OxmlElement("w:sectPr"))
        paragraph.append(properties)
        return paragraph

    def test_extracts_multiline_rows_and_detects_continuation_pages(self):
        with TemporaryDirectory() as temp_dir:
            pdf_path = Path(temp_dir) / "contents.pdf"
            self._make_contents_pdf(pdf_path)

            groups = _find_contents_page_groups(pdf_path)

        self.assertEqual([page for page, _ in groups], [0, 1])
        self.assertEqual([len(entries) for _, entries in groups], [3, 2])
        wrapped = groups[0][1][1]
        self.assertEqual(wrapped.label, "1.1.1")
        self.assertEqual(wrapped.title, "Wrapped title start continued title")
        self.assertEqual(wrapped.page, "4")

    def test_rebuilds_toc_and_preserves_page_section_markers(self):
        with TemporaryDirectory() as temp_dir:
            pdf_path = Path(temp_dir) / "contents.pdf"
            self._make_contents_pdf(pdf_path)

            document = Document()
            document.add_paragraph("TABLE OF CONTENTS")
            first_old_row = document.add_paragraph("broken page one row")
            marker_one = self._section_marker()
            first_old_row._p.addnext(marker_one)

            old_table = document.add_table(rows=1, cols=1)
            old_table.cell(0, 0).text = "duplicated table fragments"
            second_old_row = document.add_paragraph("broken page two row")
            marker_two = self._section_marker()
            second_old_row._p.addnext(marker_two)
            document.add_paragraph("INTRODUCTION")

            repaired = repair_table_of_contents(document, pdf_path)

        self.assertEqual(repaired, 5)
        self.assertEqual(len(document.tables), 0)
        self.assertEqual(
            sum(1 for node in document.element.body if node.find(".//" + qn("w:sectPr")) is not None),
            2,
        )
        paragraphs = [paragraph.text for paragraph in document.paragraphs]
        self.assertIn("1.1    Overview of the whole system\t3", paragraphs)
        self.assertIn(
            "1.1.1    Wrapped title start continued title\t4",
            paragraphs,
        )
        self.assertIn("2.1    Second page section\t8", paragraphs)
        self.assertNotIn("broken page one row", paragraphs)
        self.assertNotIn("broken page two row", paragraphs)

if __name__ == "__main__":
    unittest.main()
