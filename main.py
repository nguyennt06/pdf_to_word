"""Convert a text-based PDF to an editable DOCX document."""

from __future__ import annotations

import argparse
import re
from collections import defaultdict
from copy import deepcopy
from dataclasses import dataclass, replace
from io import BytesIO
from pathlib import Path
from tempfile import TemporaryDirectory
from uuid import uuid4

try:
    import fitz as pymupdf  # type: ignore[import-not-found]
except ImportError:
    try:
        import pymupdf  # type: ignore[import-not-found]
    except ImportError as exc:
        raise ImportError(
            "PyMuPDF is required. Install it with: pip install pymupdf"
        ) from exc

try:
    from docx import Document  # type: ignore[import-not-found]
    from docx.enum.text import WD_ALIGN_PARAGRAPH  # type: ignore[import-not-found]
    from docx.opc.constants import RELATIONSHIP_TYPE as RT  # type: ignore[import-not-found]
    from docx.opc.oxml import serialize_part_xml  # type: ignore[import-not-found]
    from docx.opc.part import Part  # type: ignore[import-not-found]
    from docx.oxml import OxmlElement, parse_xml  # type: ignore[import-not-found]
    from docx.oxml.ns import qn  # type: ignore[import-not-found]
    from docx.text.paragraph import Paragraph  # type: ignore[import-not-found]
    from docx.text.run import Run  # type: ignore[import-not-found]
except ImportError as exc:
    raise ImportError(
        "python-docx is required. Install it with: pip install python-docx"
    ) from exc

try:
    from fontTools.ttLib import TTFont  # type: ignore[import-not-found]
    from pdf2docx import Converter  # type: ignore[import-not-found]
    from pdf2docx.common import docx as pdf2docx_docx  # type: ignore[import-not-found]
    from pdf2docx.common.share import decode  # type: ignore[import-not-found]
    from pdf2docx.font.Fonts import Fonts  # type: ignore[import-not-found]
    from pdf2docx.page.RawPage import RawPage  # type: ignore[import-not-found]
    from pdf2docx.text.Lines import Lines  # type: ignore[import-not-found]
except ImportError as exc:
    raise ImportError(
        "pdf2docx is required. Install it with: pip install pdf2docx"
    ) from exc


ROMAN_HEADING_PATTERN = re.compile(
    r"^(?:\([IVXLCDM]+\)|[IVXLCDM]+(?:[.)]|\s*[-–—:]))\s+\S.*$",
    re.IGNORECASE,
)
NUMBERED_HEADING_PATTERN = re.compile(
    r"^(?P<label>\d+(?:\.\d+){0,7})(?:[.)])?"
    r"(?:\s*[-–—:]\s*|\s+)(?P<title>\S.*)$"
)
PARENTHESIZED_HEADING_PATTERN = re.compile(
    r"^\((?P<label>\d+(?:\.\d+){0,7})\)\s+(?P<title>\S.*)$"
)
CHAPTER_HEADING_PATTERN = re.compile(
    r"^(?:CHAPTER|PART|APPENDIX|CHƯƠNG|PHẦN)\s+"
    r"(?:\d+(?:\.\d+)*|[IVXLCDM]+|[A-Z])"
    r"(?:\s*[.\-–—:])?(?:\s+\S.*)?$",
    re.IGNORECASE,
)
SECTION_HEADING_PATTERN = re.compile(
    r"^(?P<kind>SECTION|SUBSECTION|MỤC|TIỂU MỤC)\s+"
    r"(?P<label>\d+(?:\.\d+)*|[IVXLCDM]+|[A-Z])"
    r"(?:\s*[.\-–—:])?(?:\s+\S.*)?$",
    re.IGNORECASE,
)
KNOWN_HEADING_PATTERN = re.compile(
    r"^(?:"
    r"ABSTRACT|ACKNOWLEDG(?:E)?MENTS?|PREFACE|FOREWORD|INTRODUCTION|"
    r"CONCLUSIONS?|REFERENCES|BIBLIOGRAPHY|TABLE OF CONTENTS|CONTENTS|"
    r"LIST OF (?:FIGURES|TABLES|ABBREVIATIONS)|APPENDIX(?:\s+.*)?|"
    r"MỤC LỤC|LỜI CẢM ƠN|LỜI NÓI ĐẦU|MỞ ĐẦU|"
    r"DANH MỤC(?:\s+.+)?|KẾT LUẬN|TÀI LIỆU THAM KHẢO|PHỤ LỤC(?:\s+.*)?"
    r")$",
    re.IGNORECASE,
)
BULLET_PREFIXES = ("•", "●", "◦", "▪", "‣", "⁃")
EMU_PER_TWIP = 635
FONT_CONTENT_TYPE = "application/vnd.openxmlformats-officedocument.obfuscatedFont"
FONT_EMBED_TAGS = {
    "regular": "embedRegular",
    "bold": "embedBold",
    "italic": "embedItalic",
    "bold_italic": "embedBoldItalic",
}


@dataclass(frozen=True)
class PdfFont:
    xref: int
    family: str
    style: str
    data: bytes
    codepoints: frozenset[int]
    subsetted: bool
    allow_subsetting: bool
    source_name: str


@dataclass(frozen=True)
class FontRecord:
    descriptor: str
    source_descriptor: str
    name: str
    line_height: float | None


def _font_name(font, *name_ids: int) -> str:
    """Read the first non-empty value from an OpenType name table."""
    if "name" not in font:
        return ""
    names = font["name"]
    for name_id in name_ids:
        value = names.getDebugName(name_id)
        if value and value.strip():
            return value.strip()
    return ""


def _remove_subset_prefix(name: str) -> str:
    """Remove the six-letter prefix used by subset PDF fonts."""
    return re.sub(r"^[A-Z]{6}\+", "", name).strip()


def _font_descriptor(name: str) -> str:
    """Normalize a font name for matching without discarding its face name."""
    return re.sub(r"[\W_]+", "", _remove_subset_prefix(name), flags=re.UNICODE).upper()


def _strip_subfamily(full_name: str, subfamily: str) -> str:
    """Remove a trailing subfamily even when its separators differ."""
    target = _font_descriptor(subfamily)
    if not target:
        return full_name.strip(" -_")

    for index in range(len(full_name)):
        if _font_descriptor(full_name[index:]) == target:
            return full_name[:index].strip(" -_")
    return full_name.strip(" -_")


def _axis_suffix(font, family: str) -> str:
    """Name non-standard weight and width axes that Word cannot select."""
    os2 = font.get("OS/2")
    if os2 is None:
        return ""

    weight_names = {
        100: "Thin",
        200: "Extra Light",
        300: "Light",
        500: "Medium",
        600: "Semi Bold",
        800: "Extra Bold",
        900: "Black",
        950: "Extra Black",
    }
    width_names = {
        1: "Ultra Condensed",
        2: "Extra Condensed",
        3: "Condensed",
        4: "Semi Condensed",
        6: "Semi Expanded",
        7: "Expanded",
        8: "Extra Expanded",
        9: "Ultra Expanded",
    }

    suffixes = []
    weight = int(getattr(os2, "usWeightClass", 400))
    width = int(getattr(os2, "usWidthClass", 5))
    if weight not in (400, 700):
        suffixes.append(weight_names.get(weight, f"Weight {weight}"))
    if width != 5:
        suffixes.append(width_names.get(width, f"Width {width}"))

    searchable_family = re.sub(r"(?<=[a-z0-9])(?=[A-Z])", " ", family)

    def already_named(suffix: str) -> bool:
        words = re.findall(r"\w+", suffix, flags=re.UNICODE)
        pattern = r"(?<!\w)" + r"[\W_]*".join(map(re.escape, words)) + r"(?!\w)"
        return re.search(pattern, searchable_family, flags=re.IGNORECASE) is not None

    return " ".join(
        suffix for suffix in suffixes if not already_named(suffix)
    )


def _read_font_family(font) -> str:
    """Recover a usable family when an embedded font omits that field."""
    family = _font_name(font, 1, 21, 16)
    subfamily = _font_name(font, 2, 22, 17)
    if not family:
        full_name = _remove_subset_prefix(_font_name(font, 4, 6))
        if not full_name:
            raise ValueError("embedded font has no usable name")
        family = _strip_subfamily(full_name, subfamily)

    family = _remove_subset_prefix(family)
    if not family:
        raise ValueError("embedded font has no usable family")

    suffix = _axis_suffix(font, family)
    return f"{family} {suffix}".strip()


def _extract_font_records(fonts_class, fitz_doc):
    """Read font metadata while retaining the PDF face name for exact matching."""
    records = []
    xrefs = sorted({item[0] for page in fitz_doc for item in page.get_fonts()})

    for xref in xrefs:
        try:
            source_name, extension, _, data = fitz_doc.extract_font(xref)
        except Exception:
            continue
        if not source_name:
            continue

        try:
            source_name = decode(source_name)
        except (TypeError, ValueError):
            source_name = str(source_name)
        family = _fallback_font_family(source_name) or f"PDF Font {xref}"
        line_height = None

        if (extension or "").lower() not in {"n/a", "cff"} and data:
            try:
                font = TTFont(BytesIO(data), lazy=False)
            except Exception:
                font = None
            if font is not None:
                try:
                    try:
                        family = _read_font_family(font)
                    except Exception:
                        pass
                    try:
                        line_height = fonts_class.get_line_height_factor(font)
                    except Exception:
                        pass
                finally:
                    font.close()

        records.append(
            FontRecord(
                descriptor=_font_descriptor(family),
                source_descriptor=_font_descriptor(source_name),
                name=family,
                line_height=line_height,
            )
        )

    return fonts_class(records)


def _find_pdf_font(fonts, font_name: str):
    """Match a PDF span to a parsed font without accepting empty names."""
    target = _font_descriptor(font_name)
    if not target:
        return None
    candidates = [
        font
        for font in fonts
        if font.descriptor or getattr(font, "source_descriptor", "")
    ]

    exact = [
        font
        for font in candidates
        if target == getattr(font, "source_descriptor", "")
    ]
    if exact:
        return min(
            exact,
            key=lambda item: (
                item.descriptor != target,
                item.line_height is not None,
                item.name.casefold(),
                item.descriptor,
            ),
        )

    exact = [font for font in candidates if target == font.descriptor]
    if exact:
        return min(exact, key=lambda item: item.name.casefold())

    contained = [font for font in candidates if font.descriptor in target]
    if contained:
        return min(
            contained,
            key=lambda item: (-len(item.descriptor), item.name.casefold()),
        )

    containing = [font for font in candidates if target in font.descriptor]
    if containing:
        return min(
            containing,
            key=lambda item: (len(item.descriptor), item.name.casefold()),
        )
    return None


def _add_hyperlink(paragraph, url: str, label: str):
    """Create a hyperlink whose visible run receives the PDF formatting."""
    relationship_id = paragraph.part.relate_to(
        url,
        RT.HYPERLINK,
        is_external=True,
    )
    hyperlink = OxmlElement("w:hyperlink")
    hyperlink.set(qn("r:id"), relationship_id)
    hyperlink.set(qn("w:history"), "1")

    run_element = OxmlElement("w:r")
    run_properties = OxmlElement("w:rPr")
    run_style = OxmlElement("w:rStyle")
    run_style.set(qn("w:val"), "Hyperlink")
    run_properties.append(run_style)
    run_element.append(run_properties)
    hyperlink.append(run_element)
    paragraph._p.append(hyperlink)

    run = Run(run_element, paragraph)
    run.text = label
    return run


_pdf2docx_create_section = RawPage._create_section
_pdf2docx_parse_line_break = Lines.parse_line_break
_pdf2docx_split_text_rows = Lines.split_vertically_by_text


def _create_pdf_section(num_columns, elements, horizontal_range, y_reference):
    """Keep all blocks when a two-column split produces extra groups."""
    if num_columns == 2 and len(elements.group_by_columns()) != 2:
        num_columns = 1
    return _pdf2docx_create_section(
        num_columns,
        elements,
        horizontal_range,
        y_reference,
    )


def _largest_font_size(row) -> float:
    """Return the largest text size in a physical PDF row."""
    return max(
        (
            getattr(span, "size", 0)
            for line in row
            for span in line.spans
        ),
        default=0,
    )


def _split_pdf_text_rows(lines, free_space_ratio, paragraph_ratio):
    """Keep visually separate PDF rows in separate paragraphs."""
    groups = _pdf2docx_split_text_rows(
        lines,
        free_space_ratio,
        paragraph_ratio,
    )
    result = []

    for group in groups:
        rows = group.group_by_physical_rows()
        current = Lines()
        previous_size = 0
        previous_row = None

        for row in rows:
            size = _largest_font_size(row)
            smaller = min(previous_size, size)
            larger = max(previous_size, size)
            gap = row.bbox[1] - previous_row.bbox[3] if previous_row else 0
            size_change = smaller and larger >= smaller * 1.5
            separate_titles = smaller >= 18 and gap >= smaller * 0.5
            if current and (size_change or separate_titles):
                result.append(current)
                current = Lines()
            current.extend(row)
            previous_size = size
            previous_row = row

        if current:
            result.append(current)

    return result


def _parse_pdf_line_breaks(
    lines,
    bounding_box,
    width_ratio,
    free_space_ratio,
):
    """Separate adjacent PDF rows when their font sizes differ sharply."""
    _pdf2docx_parse_line_break(
        lines,
        bounding_box,
        width_ratio,
        free_space_ratio,
    )

    rows = lines.group_by_physical_rows()
    for row, next_row in zip(rows, rows[1:]):
        current_size = _largest_font_size(row)
        next_size = _largest_font_size(next_row)
        smaller = min(current_size, next_size)
        larger = max(current_size, next_size)
        if smaller and larger >= smaller * 1.5:
            row[-1].line_break = 1


def configure_font_reader() -> None:
    """Use stricter font-name handling during pdf2docx conversion."""
    Fonts.get_font_family_name = staticmethod(_read_font_family)
    Fonts.extract = classmethod(_extract_font_records)
    Fonts.get = _find_pdf_font
    RawPage._create_section = staticmethod(_create_pdf_section)
    Lines.split_vertically_by_text = _split_pdf_text_rows
    Lines.parse_line_break = _parse_pdf_line_breaks
    pdf2docx_docx.add_hyperlink = _add_hyperlink


def _font_style(font, source_name: str) -> str:
    """Map an OpenType face to one of Word's four embedded-font slots."""
    subfamily = _font_name(font, 2, 22, 17).casefold()
    os2 = font.get("OS/2")
    head = font.get("head")
    selection = int(getattr(os2, "fsSelection", 0))
    weight = int(getattr(os2, "usWeightClass", 400))
    mac_style = int(getattr(head, "macStyle", 0))
    pdf_name = _remove_subset_prefix(source_name).casefold()

    bold = (
        "bold" in subfamily
        or "bold" in pdf_name
        or "black" in subfamily
        or "black" in pdf_name
        or weight >= 600
        or bool(selection & 0x20)
        or bool(mac_style & 0x01)
    )
    italic = (
        "italic" in subfamily
        or "italic" in pdf_name
        or "oblique" in subfamily
        or "oblique" in pdf_name
        or bool(selection & 0x201)
        or bool(mac_style & 0x02)
    )

    if bold and italic:
        return "bold_italic"
    if bold:
        return "bold"
    if italic:
        return "italic"
    return "regular"


def _run_style(flags: int) -> str:
    """Map PyMuPDF's bold and italic flags to a Word font slot."""
    bold = bool(flags & 2**4)
    italic = bool(flags & 2**1)
    if bold and italic:
        return "bold_italic"
    if bold:
        return "bold"
    if italic:
        return "italic"
    return "regular"


def _word_run_style(run) -> str:
    """Map a Word run to one of the embedded-font slots."""
    bold = run.bold is True
    italic = run.italic is True
    if bold and italic:
        return "bold_italic"
    if bold:
        return "bold"
    if italic:
        return "italic"
    return "regular"


def _font_embedding_issue(font, subsetted: bool) -> str | None:
    """Return why a font cannot be copied into an editable DOCX."""
    os2 = font.get("OS/2")
    if os2 is None:
        return "missing OS/2 embedding permissions"

    fs_type = int(getattr(os2, "fsType", 0))
    permission = fs_type & 0x000F
    if permission & 0x0002:
        return "restricted by the font license"
    if permission & 0x0004:
        return "licensed for preview and print only"
    if permission not in (0, 0x0008):
        return f"unsupported embedding permission 0x{permission:04X}"
    if fs_type & 0x0200:
        return "licensed for bitmap embedding only"
    if fs_type & 0x0100 and subsetted:
        return "does not allow font subsetting"
    return None


def _fallback_font_family(source_name: str) -> str:
    """Build a readable family name from a PDF font resource name."""
    name = _remove_subset_prefix(source_name)
    separated_style = re.compile(
        r"(?:[,;]|[-_])\s*(?:regular|roman|book|bold|italic|oblique|"
        r"bold[ _-]?italic|bold[ _-]?oblique)(?:MT)?$",
        re.IGNORECASE,
    )
    spaced_style = re.compile(
        r"\s+(?:regular|book|bold|italic|oblique|"
        r"bolditalic|boldoblique)(?:MT)?$",
        re.IGNORECASE,
    )
    family = separated_style.sub("", name)
    family = spaced_style.sub("", family).strip(" ,;-_") or name
    family = re.sub(r"(?:PS)?MT$", "", family, flags=re.IGNORECASE)
    return family.replace("-", " ").strip()


def _read_embeddable_fonts(pdf_path: Path) -> tuple[list[PdfFont], list[str]]:
    """Read reusable font faces already embedded in a PDF."""
    extracted = []
    warnings = []

    with pymupdf.open(pdf_path) as pdf:
        xrefs = sorted({font[0] for page in pdf for font in page.get_fonts()})
        for xref in xrefs:
            try:
                source_name, extension, _, data = pdf.extract_font(xref)
            except Exception as error:
                warnings.append(f"font xref {xref}: extraction failed ({error})")
                continue
            source_name = source_name or f"font xref {xref}"
            if (
                (extension or "").lower() not in {"ttf", "otf"}
                or not data
                or len(data) < 32
            ):
                warnings.append(f"{source_name}: unsupported or missing font data")
                continue

            try:
                font = TTFont(BytesIO(data), lazy=False)
            except Exception as error:
                warnings.append(f"{source_name}: invalid font data ({error})")
                continue

            try:
                try:
                    family = _read_font_family(font)
                except ValueError:
                    family = _fallback_font_family(source_name)
                style = _font_style(font, source_name)
                subsetted = re.match(r"^[A-Z]{6}\+", source_name) is not None
                issue = _font_embedding_issue(font, subsetted)
                codepoints = frozenset(font.getBestCmap() or {})
                os2 = font.get("OS/2")
                fs_type = int(getattr(os2, "fsType", 0))
            except Exception as error:
                warnings.append(f"{source_name}: unreadable font metadata ({error})")
                continue
            finally:
                font.close()

            if issue:
                warnings.append(f"{family}: {issue}")
                continue

            extracted.append(
                PdfFont(
                    xref=xref,
                    family=family,
                    style=style,
                    data=data,
                    codepoints=codepoints,
                    subsetted=subsetted,
                    allow_subsetting=not bool(fs_type & 0x0100),
                    source_name=source_name,
                )
            )

        usage: dict[tuple[int, str], set[int]] = defaultdict(set)
        known_xrefs = {font.xref for font in extracted}
        for page in pdf:
            xrefs_by_name: dict[str, list[int]] = defaultdict(list)
            for item in page.get_fonts():
                xref, source_name = item[0], item[3]
                if xref in known_xrefs:
                    xrefs_by_name[_font_descriptor(source_name)].append(xref)

            for block in page.get_text("dict").get("blocks", []):
                for line in block.get("lines", []):
                    for span in line.get("spans", []):
                        descriptor = _font_descriptor(span.get("font", ""))
                        codepoints = {
                            ord(character)
                            for character in span.get("text", "")
                            if character.isprintable() and not character.isspace()
                        }
                        style = _run_style(int(span.get("flags", 0)))
                        for xref in xrefs_by_name.get(descriptor, []):
                            usage[(xref, style)].update(codepoints)

    grouped: dict[tuple[str, str], list[PdfFont]] = defaultdict(list)
    for font in extracted:
        used_styles = sorted(
            style for xref, style in usage if xref == font.xref
        ) or [font.style]
        for style in used_styles:
            candidate = replace(font, style=style)
            grouped[(font.family.casefold(), style)].append(candidate)

    selected = []
    for candidates in grouped.values():
        family = candidates[0].family
        style = candidates[0].style
        required = set().union(
            *(usage[(font.xref, font.style)] for font in candidates)
        )
        unique = []
        seen_data = set()
        for candidate in candidates:
            fingerprint = hash(candidate.data)
            if fingerprint in seen_data and any(
                item.data == candidate.data for item in unique
            ):
                continue
            seen_data.add(fingerprint)
            unique.append(candidate)

        def rank(candidate: PdfFont):
            covered = len(required & candidate.codepoints)
            missing = len(required - candidate.codepoints)
            return (
                -covered,
                missing,
                -len(candidate.codepoints),
                -len(candidate.data),
                candidate.source_name.casefold(),
            )

        chosen = min(unique, key=rank)
        covered = required & chosen.codepoints
        missing = required - chosen.codepoints
        label = f"{family} {style.replace('_', ' ')}"
        if not chosen.codepoints:
            warnings.append(
                f"{label}: font has no Unicode character map and was not embedded"
            )
            continue
        if required and not covered:
            warnings.append(
                f"{label}: font covers none of its used Unicode characters and "
                "was not embedded"
            )
            continue

        selected.append(chosen)
        if len(unique) > 1:
            detail = f"; missing {len(missing)} used characters" if missing else ""
            warnings.append(
                f"{label}: multiple PDF faces share one Word slot; "
                f"using {chosen.source_name}{detail}"
            )
        elif missing:
            warnings.append(
                f"{label}: embedded font lacks {len(missing)} used Unicode "
                "characters; Word may substitute those glyphs"
            )

    return selected, warnings


def _obfuscate_font(data: bytes, font_key) -> bytes:
    """Encode a font part as required by WordprocessingML."""
    encoded = bytearray(data)
    key = font_key.bytes[::-1]
    for index in range(min(32, len(encoded))):
        encoded[index] ^= key[index % len(key)]
    return bytes(encoded)


def _font_table_entry(font_table, family: str):
    """Find or add a family in the Word font table."""
    for entry in font_table.findall(qn("w:font")):
        name = entry.get(qn("w:name"))
        if name and name.casefold() == family.casefold():
            return entry

    entry = OxmlElement("w:font")
    entry.set(qn("w:name"), family)
    font_table.append(entry)
    return entry


def _enable_embedded_fonts(document: Document, save_subsets: bool) -> None:
    """Tell Word how the embedded font data should be saved."""
    settings = document.settings._element
    if settings.find(qn("w:embedTrueTypeFonts")) is None:
        settings.append(OxmlElement("w:embedTrueTypeFonts"))

    subset_setting = settings.find(qn("w:saveSubsetFonts"))
    if save_subsets and subset_setting is None:
        settings.append(OxmlElement("w:saveSubsetFonts"))
    elif not save_subsets and subset_setting is not None:
        settings.remove(subset_setting)


def embed_pdf_fonts(
    document: Document,
    pdf_path: Path,
    heading_font_requirements: dict[tuple[str, str], set[int]] | None = None,
) -> tuple[int, list[str]]:
    """Embed licensed font faces already stored in the source PDF."""
    fonts, warnings = _read_embeddable_fonts(pdf_path)
    if not fonts:
        return 0, warnings

    try:
        font_table_part = document.part.part_related_by(RT.FONT_TABLE)
    except KeyError:
        warnings.append("DOCX package has no font table")
        return 0, warnings

    font_table = parse_xml(font_table_part.blob)
    package = document.part.package
    embedded = 0

    for font in sorted(fonts, key=lambda item: (item.family.casefold(), item.style)):
        required = (heading_font_requirements or {}).get(
            (_font_descriptor(font.family), font.style),
            set(),
        )
        missing = required - font.codepoints
        if missing:
            warnings.append(
                f"{font.family} {font.style.replace('_', ' ')}: PDF subset "
                f"lacks {len(missing)} normalized heading characters; using "
                "an installed font instead"
            )
            continue

        entry = _font_table_entry(font_table, font.family)
        tag = FONT_EMBED_TAGS[font.style]
        if entry.find(qn(f"w:{tag}")) is not None:
            continue

        font_key = uuid4()
        part_name = package.next_partname("/word/fonts/font%d.odttf")
        font_part = Part(
            part_name,
            FONT_CONTENT_TYPE,
            _obfuscate_font(font.data, font_key),
            package,
        )
        relationship_id = font_table_part.relate_to(font_part, RT.FONT)

        reference = OxmlElement(f"w:{tag}")
        reference.set(qn("r:id"), relationship_id)
        reference.set(qn("w:fontKey"), f"{{{str(font_key).upper()}}}")
        if font.subsetted:
            reference.set(qn("w:subsetted"), "true")
        entry.append(reference)
        embedded += 1

    if embedded:
        font_table_part._blob = serialize_part_xml(font_table)
        _enable_embedded_fonts(
            document,
            save_subsets=all(font.allow_subsetting for font in fonts),
        )

    return embedded, warnings


def _iter_paragraphs(document: Document):
    """Yield every paragraph in the main story, headers, and footers."""
    stories = [(document.element.body, document._body)]
    for relationship in document.part.rels.values():
        if relationship.reltype in {RT.HEADER, RT.FOOTER}:
            stories.append((relationship.target_part.element, document._body))

    seen_roots: set[object] = set()
    for root, owner in stories:
        if root in seen_roots:
            continue
        seen_roots.add(root)
        for paragraph_element in root.xpath(".//w:p"):
            yield Paragraph(paragraph_element, owner)


def _iter_runs(paragraph):
    """Yield direct runs, including runs stored inside hyperlinks."""
    for run_element in paragraph._p.xpath(".//w:r"):
        parents = run_element.xpath("ancestor::w:p[1]")
        if parents and parents[0] is paragraph._p:
            yield Run(run_element, paragraph)


def _inline_height_emu(paragraph) -> int | None:
    """Get the tallest inline image height in EMUs."""
    heights = []

    for inline in paragraph._p.xpath(".//wp:inline"):
        parents = inline.xpath("ancestor::w:p[1]")
        if not parents or parents[0] is not paragraph._p:
            continue

        extents = inline.xpath("./wp:extent")
        if not extents or not extents[0].get("cy"):
            continue

        try:
            height = int(extents[0].get("cy"))
            effects = inline.xpath("./wp:effectExtent")
            if effects:
                height += max(0, int(effects[0].get("t") or 0))
                height += max(0, int(effects[0].get("b") or 0))
        except (TypeError, ValueError):
            continue

        heights.append(height)

    return max(heights) if heights else None


def _spacing_chain(document: Document, paragraph):
    """Yield spacing settings from most to least specific."""
    paragraph_properties = paragraph._p.pPr
    if paragraph_properties is not None:
        spacing = paragraph_properties.find(qn("w:spacing"))
        if spacing is not None:
            yield spacing

    style = paragraph.style
    visited = set()
    while style is not None and id(style._element) not in visited:
        visited.add(id(style._element))
        style_properties = getattr(style._element, "pPr", None)
        if style_properties is not None:
            spacing = style_properties.find(qn("w:spacing"))
            if spacing is not None:
                yield spacing
        style = style.base_style

    yield from document.styles.element.xpath(
        "./w:docDefaults/w:pPrDefault/w:pPr/w:spacing"
    )


def _spacing_value(
    document: Document,
    paragraph,
    attribute: str,
) -> str | None:
    """Get a spacing value after style inheritance."""
    key = qn(f"w:{attribute}")
    for spacing in _spacing_chain(document, paragraph):
        value = spacing.get(key)
        if value is not None:
            return value
    return None


def fix_inline_image_spacing(document: Document) -> int:
    """Let lines expand when exact spacing would clip an inline image."""
    count = 0

    for paragraph in _iter_paragraphs(document):
        image_height = _inline_height_emu(paragraph)
        if image_height is None:
            continue

        line_rule = _spacing_value(document, paragraph, "lineRule")
        line = _spacing_value(document, paragraph, "line")
        if line_rule != "exact" or line is None:
            continue

        try:
            line_twips = int(line)
        except ValueError:
            continue
        if line_twips * EMU_PER_TWIP >= image_height:
            continue

        spacing = paragraph._p.get_or_add_pPr().get_or_add_spacing()
        spacing.set(qn("w:line"), str(line_twips))
        spacing.set(qn("w:lineRule"), "atLeast")
        count += 1

    return count


def fix_bullet_fonts(document: Document) -> int:
    """Replace Symbol bullets that Word renders as squares."""
    count = 0

    for paragraph in _iter_paragraphs(document):
        runs = list(_iter_runs(paragraph))
        for index, run in enumerate(runs):
            if run.text.strip() not in {"•", "●", "◦"}:
                continue
            if (run.font.name or "").lower() not in {"symbol", "symbolmt"}:
                continue

            run.font.name = "Times New Roman"
            rfonts = run._element.get_or_add_rPr().rFonts
            for name in ("ascii", "hAnsi", "eastAsia", "cs"):
                rfonts.set(qn(f"w:{name}"), "Times New Roman")
            if index + 1 < len(runs) and not runs[index + 1].text.startswith(
                (" ", "\t")
            ):
                run.text = f"{run.text} "
            count += 1

    return count


def complete_font_settings(document: Document) -> int:
    """Complete compatible Word font fields without crossing script families."""
    count = 0

    for paragraph in _iter_paragraphs(document):
        for run in _iter_runs(paragraph):
            run_properties = run._r.rPr
            if run_properties is None:
                continue

            changed = False
            fonts = run_properties.find(qn("w:rFonts"))
            if fonts is not None:
                ascii_font = fonts.get(qn("w:ascii"))
                high_ansi_font = fonts.get(qn("w:hAnsi"))

                if (
                    ascii_font
                    and not high_ansi_font
                    and not fonts.get(qn("w:hAnsiTheme"))
                ):
                    fonts.set(qn("w:hAnsi"), ascii_font)
                    high_ansi_font = ascii_font
                    changed = True
                elif (
                    high_ansi_font
                    and not ascii_font
                    and not fonts.get(qn("w:asciiTheme"))
                ):
                    fonts.set(qn("w:ascii"), high_ansi_font)
                    ascii_font = high_ansi_font
                    changed = True

                latin_font = high_ansi_font or ascii_font
                if (
                    latin_font
                    and not fonts.get(qn("w:cs"))
                    and not fonts.get(qn("w:cstheme"))
                ):
                    fonts.set(qn("w:cs"), latin_font)
                    changed = True

            size = run_properties.find(qn("w:sz"))
            complex_size = run_properties.find(qn("w:szCs"))
            if size is not None and size.get(qn("w:val")):
                if complex_size is None:
                    complex_size = OxmlElement("w:szCs")
                    run_properties.insert(
                        run_properties.index(size) + 1,
                        complex_size,
                    )
                if not complex_size.get(qn("w:val")):
                    complex_size.set(qn("w:val"), size.get(qn("w:val")))
                    changed = True

            count += changed

    return count


def normalize_text(value: str) -> str:
    """Normalize whitespace from PDF text."""
    return " ".join(value.split())


def detect_heading_scheme(document: Document) -> tuple[bool, bool]:
    """Check whether the document uses Roman or chapter headings."""
    has_roman_sections = False
    has_chapter_labels = False

    for paragraph in document.paragraphs:
        samples = (paragraph.text, *(run.text for run in _iter_runs(paragraph)))
        for sample in samples:
            text = normalize_text(sample)
            if not text:
                continue
            has_roman_sections |= ROMAN_HEADING_PATTERN.match(text) is not None
            has_chapter_labels |= CHAPTER_HEADING_PATTERN.match(text) is not None

    return has_roman_sections, has_chapter_labels


def strict_heading_style(
    text: str,
    *,
    has_roman_sections: bool,
    has_chapter_labels: bool,
) -> str | None:
    """Match a heading label to a Word Heading style."""
    if not text or len(text) > 240 or text.startswith(BULLET_PREFIXES):
        return None

    if ROMAN_HEADING_PATTERN.match(text):
        return "Heading 1"

    if CHAPTER_HEADING_PATTERN.match(text):
        return "Heading 2" if has_roman_sections else "Heading 1"

    section_match = SECTION_HEADING_PATTERN.match(text)
    if section_match:
        label = section_match.group("label")
        depth = len(label.split(".")) if label[0].isdigit() else 1
        if has_roman_sections:
            level = depth + 1
        elif has_chapter_labels:
            level = max(2, depth)
        else:
            level = depth
        if section_match.group("kind").upper() in {"SUBSECTION", "TIỂU MỤC"}:
            level = max(level, 3)
        return f"Heading {min(level, 9)}"

    if KNOWN_HEADING_PATTERN.match(text):
        return "Heading 2" if has_roman_sections else "Heading 1"

    match = PARENTHESIZED_HEADING_PATTERN.match(text)
    if match is None:
        match = NUMBERED_HEADING_PATTERN.match(text)
    if match is None:
        return None

    depth = len(match.group("label").split("."))
    if has_roman_sections:
        level = depth + 1
    elif has_chapter_labels:
        level = max(2, depth)
    else:
        level = depth

    return f"Heading {min(level, 9)}"


def paragraph_has_numbering(paragraph) -> bool:
    """Check whether Word already treats this paragraph as a list."""
    paragraph_properties = paragraph._p.pPr
    if (
        paragraph_properties is not None
        and paragraph_properties.find(qn("w:numPr")) is not None
    ):
        return True

    style = paragraph.style
    visited = set()
    while style is not None and id(style._element) not in visited:
        visited.add(id(style._element))
        style_properties = getattr(style._element, "pPr", None)
        if (
            style_properties is not None
            and style_properties.find(qn("w:numPr")) is not None
        ):
            return True
        style = style.base_style

    return False


def visual_heading_style(
    paragraph,
    text: str,
    *,
    has_roman_sections: bool,
) -> str | None:
    """Detect an unnumbered heading from its formatting."""
    is_centered_or_right = paragraph.alignment in {
        WD_ALIGN_PARAGRAPH.CENTER,
        WD_ALIGN_PARAGRAPH.RIGHT,
    }
    if (
        len(text) > 140
        or len(text.split()) > 16
        or text.startswith(BULLET_PREFIXES)
        or text.endswith((".", ",", ";", "?", "!"))
        or is_centered_or_right
    ):
        return None

    visible_runs = [run for run in _iter_runs(paragraph) if run.text.strip()]
    character_count = sum(len(run.text.strip()) for run in visible_runs)
    if character_count == 0:
        return None

    bold_count = sum(
        len(run.text.strip()) for run in visible_runs if run.bold is True
    )
    if bold_count / character_count < 0.8:
        return None

    spacing = paragraph.paragraph_format.space_before
    if spacing is None and paragraph.style is not None:
        spacing = paragraph.style.paragraph_format.space_before
    spacing_pt = getattr(spacing, "pt", 0) or 0
    if spacing_pt < 6:
        return None

    return "Heading 3" if has_roman_sections else "Heading 2"


def is_soft_break_only(run) -> bool:
    """Check whether a run contains only soft line breaks."""
    has_break = False
    for child in run._r:
        if child.tag == qn("w:rPr"):
            continue
        if child.tag != qn("w:br"):
            return False
        break_type = child.get(qn("w:type"))
        if break_type not in (None, "textWrapping"):
            return False
        has_break = True
    return has_break


def clone_split_paragraph_properties(paragraph, index: int, total: int):
    """Copy paragraph properties without repeating page or section breaks."""
    if paragraph._p.pPr is None:
        return None

    properties = deepcopy(paragraph._p.pPr)
    if index > 0:
        spacing = properties.find(qn("w:spacing"))
        if spacing is None:
            spacing = OxmlElement("w:spacing")
            properties.append(spacing)
        spacing.set(qn("w:before"), "0")

        page_break = properties.find(qn("w:pageBreakBefore"))
        if page_break is not None:
            properties.remove(page_break)

    if index < total - 1:
        section = properties.find(qn("w:sectPr"))
        if section is not None:
            properties.remove(section)

    return properties


def split_compound_heading_paragraphs(
    document: Document,
    *,
    has_roman_sections: bool,
    has_chapter_labels: bool,
) -> int:
    """Split PDF text blocks that contain several headings."""
    split_count = 0

    for paragraph in list(document.paragraphs):
        runs = paragraph.runs
        if len(paragraph._p.xpath(".//w:br")) == 0 or len(runs) < 3:
            continue
        if paragraph._p.xpath(".//w:drawing|.//w:pict"):
            continue
        if any(
            child.tag not in {qn("w:pPr"), qn("w:r")}
            for child in paragraph._p
        ):
            continue

        nonempty_indices = [
            index for index, run in enumerate(runs) if normalize_text(run.text)
        ]
        if not nonempty_indices:
            continue

        heading_starts = [
            index
            for index in nonempty_indices
            if strict_heading_style(
                normalize_text(runs[index].text),
                has_roman_sections=has_roman_sections,
                has_chapter_labels=has_chapter_labels,
            )
            is not None
        ]
        if len(heading_starts) < 2 or heading_starts[0] != nonempty_indices[0]:
            continue

        segments = [
            runs[start : heading_starts[position + 1]]
            if position + 1 < len(heading_starts)
            else runs[start:]
            for position, start in enumerate(heading_starts)
        ]

        parent = paragraph._p.getparent()
        insertion_index = parent.index(paragraph._p)
        for index, segment in enumerate(segments):
            new_paragraph = OxmlElement("w:p")
            properties = clone_split_paragraph_properties(
                paragraph, index, len(segments)
            )
            if properties is not None:
                new_paragraph.append(properties)

            for run_index, run in enumerate(segment):
                is_trailing_soft_break = (
                    run_index == len(segment) - 1 and is_soft_break_only(run)
                )
                if not is_trailing_soft_break:
                    new_paragraph.append(deepcopy(run._r))

            parent.insert(insertion_index + index, new_paragraph)

        parent.remove(paragraph._p)
        split_count += 1

    return split_count


def style_color_element(style):
    """Find the color inherited from a style."""
    visited = set()
    while style is not None and id(style._element) not in visited:
        visited.add(id(style._element))
        run_properties = getattr(style._element, "rPr", None)
        if run_properties is not None:
            color = run_properties.find(qn("w:color"))
            if color is not None:
                return deepcopy(color)
        style = style.base_style
    return None


def style_font_element(style):
    """Find the font mapping inherited from a style."""
    visited = set()
    while style is not None and id(style._element) not in visited:
        visited.add(id(style._element))
        run_properties = getattr(style._element, "rPr", None)
        if run_properties is not None:
            fonts = run_properties.find(qn("w:rFonts"))
            if fonts is not None and any(fonts.attrib.values()):
                return deepcopy(fonts)
        style = style.base_style
    return None


def effective_run_color(document: Document, paragraph, run):
    """Get a run's current color."""
    run_properties = run._r.rPr
    if run_properties is not None:
        color = run_properties.find(qn("w:color"))
        if color is not None:
            return deepcopy(color)

    for style in (run.style, paragraph.style):
        color = style_color_element(style)
        if color is not None:
            return color

    default_colors = document.styles.element.xpath(
        "./w:docDefaults/w:rPrDefault/w:rPr/w:color"
    )
    if default_colors:
        return deepcopy(default_colors[0])

    color = OxmlElement("w:color")
    color.set(qn("w:val"), "auto")
    return color


def effective_run_fonts(document: Document, paragraph, run):
    """Get a run's current font mapping."""
    run_properties = run._r.rPr
    if run_properties is not None:
        fonts = run_properties.find(qn("w:rFonts"))
        if fonts is not None and any(fonts.attrib.values()):
            return deepcopy(fonts)

    for style in (run.style, paragraph.style):
        fonts = style_font_element(style)
        if fonts is not None:
            return fonts

    defaults = document.styles.element.xpath(
        "./w:docDefaults/w:rPrDefault/w:rPr/w:rFonts"
    )
    return deepcopy(defaults[0]) if defaults else None


def _is_latin_fallback_fragment(text: str) -> bool:
    """Recognize short runs made from extended Latin letters."""
    characters = [character for character in text if not character.isspace()]
    return 0 < len(characters) <= 3 and all(
        character.isalpha()
        and (
            0x00C0 <= ord(character) <= 0x024F
            or 0x1E00 <= ord(character) <= 0x1EFF
        )
        for character in characters
    )


def normalize_heading_fallback_fonts(
    paragraph,
    required_fonts: dict[tuple[str, str], set[int]] | None = None,
) -> int:
    """Replace short fallback runs with a heading's dominant font."""
    runs = []
    totals: dict[str, int] = defaultdict(int)

    for run in _iter_runs(paragraph):
        visible_text = "".join(
            character for character in run.text if not character.isspace()
        )
        if not visible_text:
            continue

        properties = run._r.rPr
        fonts = (
            properties.find(qn("w:rFonts"))
            if properties is not None
            else None
        )
        if fonts is None:
            continue

        family = fonts.get(qn("w:hAnsi")) or fonts.get(qn("w:ascii"))
        key = _font_descriptor(family or "")
        if not key:
            continue

        totals[key] += len(visible_text)
        runs.append((run, visible_text, key, deepcopy(fonts)))

    if len(totals) < 2:
        return 0

    dominant_key, dominant_count = min(
        totals.items(),
        key=lambda item: (-item[1], item[0]),
    )
    if dominant_count * 3 < sum(totals.values()) * 2:
        return 0

    dominant_fonts = max(
        (item for item in runs if item[2] == dominant_key),
        key=lambda item: len(item[1]),
    )[3]

    changed = 0
    for index, (run, visible_text, key, _) in enumerate(runs):
        if key == dominant_key or not _is_latin_fallback_fragment(visible_text):
            continue

        touches_dominant = (
            index > 0 and runs[index - 1][2] == dominant_key
        ) or (
            index + 1 < len(runs) and runs[index + 1][2] == dominant_key
        )
        if not touches_dominant:
            continue

        target = run._r.get_or_add_rPr().get_or_add_rFonts()
        target.attrib.clear()
        target.attrib.update(dominant_fonts.attrib)
        if required_fonts is not None:
            family = target.get(qn("w:hAnsi")) or target.get(qn("w:ascii"))
            required_fonts.setdefault(
                (_font_descriptor(family or ""), _word_run_style(run)),
                set(),
            ).update(ord(character) for character in visible_text)
        changed += 1

    return changed


def _is_enabled(element) -> bool:
    """Read an OOXML on/off value."""
    value = element.get(qn("w:val"))
    return value is None or value.casefold() not in {"0", "false", "off", "no"}


def _paragraph_setting(document: Document, paragraph, name: str) -> bool:
    """Read a pagination setting before changing the paragraph style."""
    properties = paragraph._p.pPr
    if properties is not None:
        element = properties.find(qn(f"w:{name}"))
        if element is not None:
            return _is_enabled(element)

    style = paragraph.style
    visited = set()
    while style is not None and id(style._element) not in visited:
        visited.add(id(style._element))
        properties = getattr(style._element, "pPr", None)
        if properties is not None:
            element = properties.find(qn(f"w:{name}"))
            if element is not None:
                return _is_enabled(element)
        style = style.base_style

    defaults = document.styles.element.xpath(
        f"./w:docDefaults/w:pPrDefault/w:pPr/w:{name}"
    )
    return _is_enabled(defaults[0]) if defaults else False


def _set_paragraph_setting(paragraph, name: str, enabled: bool) -> None:
    """Set a paragraph pagination option explicitly."""
    properties = paragraph._p.get_or_add_pPr()
    element = properties.find(qn(f"w:{name}"))
    if element is None:
        element = OxmlElement(f"w:{name}")
        properties.append(element)
    element.set(qn("w:val"), "1" if enabled else "0")


def assign_heading_style(
    document: Document,
    paragraph,
    style_name: str,
    required_fonts: dict[tuple[str, str], set[int]] | None = None,
) -> None:
    """Apply a Heading style while preserving its visible formatting."""
    pagination = {
        name: _paragraph_setting(document, paragraph, name)
        for name in ("keepNext", "keepLines", "pageBreakBefore")
    }
    original_formatting = [
        (
            run,
            effective_run_color(document, paragraph, run),
            effective_run_fonts(document, paragraph, run),
        )
        for run in _iter_runs(paragraph)
    ]
    paragraph.style = style_name
    for name, enabled in pagination.items():
        _set_paragraph_setting(paragraph, name, enabled)

    for run, original_color, original_fonts in original_formatting:
        target_color = run._r.get_or_add_rPr().get_or_add_color()
        target_color.attrib.clear()
        target_color.attrib.update(original_color.attrib)
        if original_fonts is not None:
            target_fonts = run._r.get_or_add_rPr().get_or_add_rFonts()
            target_fonts.attrib.clear()
            target_fonts.attrib.update(original_fonts.attrib)

    normalize_heading_fallback_fonts(paragraph, required_fonts)


def detect_headings(
    document: Document,
    required_fonts: dict[tuple[str, str], set[int]] | None = None,
) -> int:
    """Apply Word Heading styles to detected headings."""
    has_roman_sections, has_chapter_labels = detect_heading_scheme(document)
    split_compound_heading_paragraphs(
        document,
        has_roman_sections=has_roman_sections,
        has_chapter_labels=has_chapter_labels,
    )
    count = 0

    # Skip table cells because their numbering is usually data.
    for paragraph in document.paragraphs:
        text = normalize_text(paragraph.text)
        physical_lines = [
            line for line in paragraph.text.splitlines() if line.strip()
        ]
        if not text or len(text) > 240 or paragraph_has_numbering(paragraph):
            continue

        style_name = strict_heading_style(
            text,
            has_roman_sections=has_roman_sections,
            has_chapter_labels=has_chapter_labels,
        )
        if style_name is None:
            if len(physical_lines) > 2:
                continue
            style_name = visual_heading_style(
                paragraph,
                text,
                has_roman_sections=has_roman_sections,
            )
        if style_name is None:
            continue

        assign_heading_style(document, paragraph, style_name, required_fonts)
        count += 1

    return count


def inspect_pdf(pdf_path: Path) -> tuple[int, int]:
    """Open the PDF and count pages and embedded images."""
    with pymupdf.open(pdf_path) as pdf:
        if pdf.needs_pass:
            raise ValueError("PDF is password-protected")
        return pdf.page_count, sum(len(page.get_images(full=True)) for page in pdf)


def convert_pdf(
    pdf_path: Path,
    docx_path: Path,
    *,
    headings: bool = False,
    embed_fonts: bool = True,
) -> tuple[int, int, int, int, int, int, int, list[str]]:
    """Convert a PDF and post-process the generated DOCX."""
    page_count, image_count = inspect_pdf(pdf_path)

    with TemporaryDirectory(prefix="pdf_to_word_") as temp_dir:
        raw_docx = Path(temp_dir) / "converted.docx"
        configure_font_reader()
        converter = Converter(str(pdf_path))
        try:
            converter.convert(str(raw_docx))
        finally:
            converter.close()

        document = Document(raw_docx)
        image_fix_count = fix_inline_image_spacing(document)

        bullet_count = fix_bullet_fonts(document)
        heading_font_requirements: dict[tuple[str, str], set[int]] = {}
        heading_count = (
            detect_headings(document, heading_font_requirements)
            if headings
            else 0
        )
        font_fix_count = complete_font_settings(document)
        if embed_fonts:
            embedded_font_count, font_warnings = embed_pdf_fonts(
                document,
                pdf_path,
                heading_font_requirements,
            )
        else:
            embedded_font_count, font_warnings = 0, []
        document.save(docx_path)

    return (
        page_count,
        image_count,
        image_fix_count,
        bullet_count,
        heading_count,
        font_fix_count,
        embedded_font_count,
        font_warnings,
    )


def parse_args() -> argparse.Namespace:
    parser = argparse.ArgumentParser(
        description="Convert a text-based PDF to an editable DOCX file."
    )
    parser.add_argument("input", type=Path, help="Input PDF path")
    parser.add_argument("-o", "--output", type=Path, help="Output DOCX path")
    parser.add_argument(
        "--headings",
        action="store_true",
        help="Detect numbered headings and assign Word Heading styles",
    )
    parser.add_argument(
        "--no-embed-fonts",
        action="store_true",
        help="Keep font names but do not embed licensed PDF font subsets",
    )
    parser.add_argument("--force", action="store_true", help="Overwrite output")
    return parser.parse_args()


def main() -> int:
    args = parse_args()
    pdf_path = args.input.expanduser().resolve()
    docx_path = (args.output or pdf_path.with_suffix(".docx")).expanduser().resolve()

    if not pdf_path.is_file():
        raise SystemExit(f"Input file not found: {pdf_path}")
    if pdf_path.suffix.lower() != ".pdf":
        raise SystemExit("Input file must have a .pdf extension")
    if docx_path.exists() and not args.force:
        raise SystemExit(f"Output already exists: {docx_path} (use --force)")

    docx_path.parent.mkdir(parents=True, exist_ok=True)
    (
        pages,
        images,
        fixed,
        bullets,
        headings,
        font_fixes,
        fonts,
        font_warnings,
    ) = convert_pdf(
        pdf_path,
        docx_path,
        headings=args.headings,
        embed_fonts=not args.no_embed_fonts,
    )

    print(f"Created: {docx_path}")
    print(
        f"Pages: {pages} | Images: {images} | "
        f"Image fixes: {fixed} | Bullet fixes: {bullets} | "
        f"Headings: {headings} | Font fixes: {font_fixes} | "
        f"Embedded fonts: {fonts}"
    )
    for warning in font_warnings:
        print(f"Font warning: {warning}")
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
